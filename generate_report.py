"""
Generates CEP_Report.docx — the final submission document for
CEP Assignment 3 & 4: Event-Driven Traffic Alert System.
"""

from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
import os

doc = Document()

# ── Global style defaults ─────────────────────────────────────────────
style = doc.styles["Normal"]
font = style.font
font.name = "Times New Roman"
font.size = Pt(12)
style.paragraph_format.line_spacing = 1.5

# ── Helper functions ──────────────────────────────────────────────────

def add_title_text(text, size=28, bold=True, color=None, spacing_after=6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(spacing_after)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Times New Roman"
    if color:
        run.font.color.rgb = RGBColor(*color)
    return p


def add_heading_styled(text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = "Times New Roman"
        run.font.color.rgb = RGBColor(0, 0, 0)
    return h


def add_body(text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(6)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for run in p.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
    return p


def add_code_block(code_text, label=None):
    if label:
        p = doc.add_paragraph()
        run = p.add_run(label)
        run.bold = True
        run.font.size = Pt(11)
        run.font.name = "Times New Roman"
        p.paragraph_format.space_after = Pt(4)

    for line in code_text.strip().split("\n"):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(30, 30, 30)
        # Light gray shading
        shading = run._element.get_or_add_rPr()
        shd = shading.makeelement(qn("w:shd"), {
            qn("w:val"): "clear",
            qn("w:color"): "auto",
            qn("w:fill"): "F2F2F2",
        })
        shading.append(shd)


def add_table(headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # Header row
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.bold = True
                run.font.name = "Times New Roman"
                run.font.size = Pt(11)
        shading = cell._element.get_or_add_tcPr()
        shd = shading.makeelement(qn("w:shd"), {
            qn("w:val"): "clear",
            qn("w:color"): "auto",
            qn("w:fill"): "D9E2F3",
        })
        shading.append(shd)
    # Data rows
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = str(val)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(11)
    doc.add_paragraph()  # spacing


def read_source(filename):
    base = os.path.dirname(os.path.abspath(__file__))
    path = os.path.join(base, filename)
    with open(path, "r", encoding="utf-8") as f:
        return f.read()


# ══════════════════════════════════════════════════════════════════════
# PAGE 1 — TITLE PAGE
# ══════════════════════════════════════════════════════════════════════

# Add spacing to push content to center
for _ in range(4):
    doc.add_paragraph()

add_title_text("COMSATS University Islamabad", size=22, color=(0, 51, 102), spacing_after=12)
add_title_text("Department of Computer Science", size=16, bold=False, color=(80, 80, 80), spacing_after=30)

# Horizontal rule
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("─" * 60)
run.font.color.rgb = RGBColor(0, 51, 102)
run.font.size = Pt(10)

add_title_text("Software Design and Architecture", size=20, color=(0, 51, 102), spacing_after=8)
add_title_text("CEP Assignment 3 & 4", size=24, color=(0, 0, 0), spacing_after=8)
add_title_text("Event-Driven Traffic Alert System", size=18, bold=False, color=(60, 60, 60), spacing_after=20)

# Horizontal rule
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("─" * 60)
run.font.color.rgb = RGBColor(0, 51, 102)
run.font.size = Pt(10)

# Details table
info_data = [
    ("Submitted to:", "Dr. Nida Adnan"),
    ("Class:", "BSSE 4A"),
    ("Total Marks:", "60"),
]

for label, value in info_data:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    run_label = p.add_run(label + "  ")
    run_label.bold = True
    run_label.font.size = Pt(14)
    run_label.font.name = "Times New Roman"
    run_label.font.color.rgb = RGBColor(0, 51, 102)
    run_value = p.add_run(value)
    run_value.font.size = Pt(14)
    run_value.font.name = "Times New Roman"

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════════
# INTRODUCTION
# ══════════════════════════════════════════════════════════════════════

add_heading_styled("Introduction", level=1)

add_body(
    "This report presents the design and implementation of an Event-Driven Traffic Alert System "
    "developed in Python 3. The system simulates a metropolitan traffic monitoring infrastructure "
    "in which cameras at intersections detect vehicles, identify speed violations, report congestion, "
    "and signal traffic clearance. Events are published through a central EventBus and consumed by "
    "specialised subscriber services including AlertService, LoggingService, DashboardService, and "
    "ReportingService."
)

add_body(
    "The architecture employs three core design patterns. The Observer Pattern decouples event "
    "producers from consumers, allowing new subscribers to be added without modifying the EventBus. "
    "The Event Envelope Pattern wraps each event payload with metadata—a unique event ID, correlation "
    "ID, schema version, source identifier, and timestamp—enabling consistent tracking and routing "
    "across the system. The Idempotent Receiver Pattern prevents duplicate event processing by "
    "maintaining a set of previously seen event IDs, ensuring that network retries or duplicate "
    "deliveries do not cause repeated side effects such as issuing the same penalty twice."
)

add_body(
    "The report is structured as follows: Section CLO 3 presents the implementation code for each "
    "system component with explanations. The UML Diagram section provides a class diagram illustrating "
    "the Observer Pattern relationships. Section CLO 4 analyses three real-world scenarios—schema "
    "evolution, event flooding, and the dual write problem—evaluating architectural trade-offs and "
    "proposing solutions grounded in established software engineering principles."
)

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════════
# CLO 3 — DESIGN PATTERNS: CODE
# ══════════════════════════════════════════════════════════════════════

add_heading_styled("CLO 3 — Design Patterns: Implementation", level=1)

add_body(
    "This section presents the complete source code for each component of the Event-Driven Traffic "
    "Alert System. Each file is documented with inline comments and docstrings explaining its role "
    "within the overall architecture."
)

# ── events.py ─────────────────────────────────────────────────────────
add_heading_styled("Task 1: Event Definitions — events.py", level=2)
add_body(
    "This module defines the four event dataclasses that represent domain occurrences in the traffic "
    "monitoring system. Each event is a simple data container decorated with @dataclass for automatic "
    "constructor generation and clean string representation."
)
add_code_block(read_source("core/events.py"))

# ── event_envelope.py ─────────────────────────────────────────────────
add_heading_styled("Task 2: Event Envelope — event_envelope.py", level=2)
add_body(
    "The EventEnvelope class implements the Event Envelope Pattern. It wraps an event payload with "
    "metadata fields: a unique event_id (UUID), a correlation_id for tracing related events across "
    "a journey, a schema_version for evolution, a source_id identifying the originating camera, "
    "and a timestamp. This metadata enables routing, deduplication, and auditing."
)
add_code_block(read_source("core/event_envelope.py"))

# ── idempotent_base.py ────────────────────────────────────────────────
add_heading_styled("Task 3: Idempotent Receiver — idempotent_base.py", level=2)
add_body(
    "The IdempotentReceiver abstract base class implements the Idempotent Receiver Pattern. It "
    "maintains a seen_ids set and checks each incoming envelope's event_id before processing. "
    "If the event has already been seen, it is rejected with a duplicate notice. Subclasses "
    "implement the abstract process_event() method to define their specific handling logic."
)
add_code_block(read_source("core/idempotent_base.py"))

# ── event_bus.py ──────────────────────────────────────────────────────
add_heading_styled("Task 4: Event Bus — event_bus.py", level=2)
add_body(
    "The EventBus class implements the Observer Pattern. It maintains a list of subscribers and "
    "provides subscribe(), unsubscribe(), and publish() methods. When an event is published, the "
    "bus iterates over all registered subscribers and delivers the envelope to each via on_event(). "
    "The bus is fully decoupled—it knows nothing about event types or subscriber implementations."
)
add_code_block(read_source("core/event_bus.py"))

# ── subscribers.py ────────────────────────────────────────────────────
add_heading_styled("Task 5: Subscriber Services — subscribers.py", level=2)
add_body(
    "This module defines five classes. IEventSubscriber is the abstract interface that all subscribers "
    "implement. AlertService and LoggingService extend IdempotentReceiver for duplicate protection. "
    "DashboardService handles real-time display updates. ReportingService collects data for periodic "
    "reporting. Each service processes only the event types relevant to its domain responsibility."
)
add_code_block(read_source("core/subscribers.py"))

# ── main.py ───────────────────────────────────────────────────────────
add_heading_styled("Task 6: System Demonstration — main.py", level=2)
add_body(
    "The main.py script demonstrates the complete system in action. It creates an EventBus, "
    "instantiates all four subscriber services, subscribes them, and publishes one event of each "
    "type. It then demonstrates the unsubscribe capability by removing DashboardService and "
    "publishing an additional event to show that only the remaining subscribers receive it. "
    "A final summary displays the penalty count, log entries, and report data points collected."
)
add_code_block(read_source("main.py"))

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════════
# UML DIAGRAM
# ══════════════════════════════════════════════════════════════════════

add_heading_styled("UML Class Diagram — Observer Pattern", level=1)

add_body(
    "The following UML class diagram illustrates the Observer Pattern as implemented in the traffic "
    "alert system. It shows the relationships between IEventSubscriber (abstract interface), "
    "IdempotentReceiver (abstract base with deduplication), EventBus (the subject/publisher), "
    "and the four concrete subscriber services. Dashed arrows indicate interface implementation, "
    "solid arrows indicate inheritance, and the association arrow from EventBus to IEventSubscriber "
    "represents the subscribers relationship with multiplicity *."
)

# Check if UML image exists
uml_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "report", "UML_Diagram.png")
if os.path.exists(uml_path):
    doc.add_picture(uml_path, width=Inches(6.0))
    last_paragraph = doc.paragraphs[-1]
    last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
else:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(20)
    p.paragraph_format.space_after = Pt(20)
    run = p.add_run("[INSERT UML_Diagram.png HERE]")
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(180, 0, 0)
    run.font.name = "Times New Roman"

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = note.add_run(
        "Export your draw.io diagram as UML_Diagram.png and place it in the report/ folder.\n"
        "Then re-run this script to embed it automatically."
    )
    run2.font.size = Pt(10)
    run2.font.italic = True
    run2.font.color.rgb = RGBColor(100, 100, 100)

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════════
# CLO 4 — SCENARIO 1: SCHEMA EVOLUTION
# ══════════════════════════════════════════════════════════════════════

add_heading_styled("CLO 4 — Scenario 1: Schema Evolution", level=1)

add_heading_styled("Context", level=2)
add_body(
    "The traffic alert system has been operational for six months. The city authority now requires "
    "a new field called lane_number to be added to VehicleDetectedEvent. Approximately 200 subscriber "
    "instances are deployed in production, consuming events under the original schema. Any modification "
    "risks deserialization failures, data inconsistencies, or silent processing errors."
)

add_heading_styled("Option A — Backward Compatibility (Optional Field with Default)", level=2)
add_body(
    "Under this approach, the existing VehicleDetectedEvent dataclass is modified in place. The new "
    "field is declared as lane_number: int = None. Because the field carries a default value, all "
    "existing code that constructs a VehicleDetectedEvent without supplying lane_number will continue "
    "to function. Subscribers that do not reference lane_number will simply ignore it."
)
add_body(
    "Advantages: Zero disruption to running subscribers — the 200 deployed instances need no "
    "redeployment or code change. There is a single source of truth with one event class to maintain "
    "and test. Minimal operational overhead with no version-routing logic needed in the EventBus."
)
add_body(
    "Risks: Semantic ambiguity of None — subscribers cannot distinguish between \"the producer "
    "genuinely does not know the lane\" and \"the producer is running old code.\" Accumulated technical "
    "debt as repeated optional fields erode type safety. No compile-time enforcement in Python's "
    "dynamic type system."
)

add_heading_styled("Option B — Schema Versioning (New Event Class)", level=2)
add_body(
    "Under this approach, the original VehicleDetectedEvent is left untouched. A new class "
    "VehicleDetectedEvent_v2 is introduced with lane_number as a required field, and schema_version "
    "is set to 2 in the EventEnvelope. Producers emit v2 payloads; subscribers that only understand "
    "v1 must be updated or they will silently ignore v2 events."
)
add_body(
    "Advantages: Explicit contract boundary — each version has a precise, immutable schema. "
    "Safe parallel operation with version-aware routing possible."
)
add_body(
    "Risks: Subscriber migration burden — all 200 instances must eventually be updated. Class "
    "proliferation over successive schema changes (_v3, _v4...). Coordinated deployment required "
    "across teams and infrastructure tiers."
)

add_heading_styled("Architecture Decision Record (ADR)", level=2)

add_heading_styled("Problem", level=3)
add_body(
    "The system must incorporate a lane_number field into VehicleDetectedEvent to satisfy a new "
    "regulatory requirement. 200 subscriber instances consume events under the original schema. "
    "Modifying the schema risks breaking subscribers, causing data loss, or creating operational "
    "blind spots in the traffic monitoring pipeline."
)

add_heading_styled("Decision", level=3)
add_body(
    "We adopt Option A — Backward Compatibility, adding lane_number as an optional field with a "
    "default value of None on the existing VehicleDetectedEvent dataclass. The rationale is threefold: "
    "(1) Operational continuity is the highest priority with 200 live instances. (2) The change is "
    "additive, not destructive, satisfying the Open/Closed Principle. (3) Version proliferation is "
    "premature — the system has undergone only one schema change in six months."
)

add_heading_styled("Consequences", level=3)
add_body(
    "Gains: Immediate backward compatibility with all 200 subscribers; a single event class to "
    "maintain; no changes to EventBus routing logic; faster time-to-delivery for the city authority."
)
add_body(
    "Losses: Reduced type safety on the lane_number field; no clear versioning boundary for future "
    "developers; if multiple additive changes accumulate, migration to explicit versioning will carry "
    "a larger refactoring cost."
)

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════════
# CLO 4 — SCENARIO 2: EVENT FLOODING
# ══════════════════════════════════════════════════════════════════════

add_heading_styled("CLO 4 — Scenario 2: Event Flooding", level=1)

add_heading_styled("Context", level=2)
add_body(
    "During a football match, all twelve intersections near the stadium generate events simultaneously. "
    "The EventBus receives 500 events per second. DashboardService can process only 80 events per "
    "second. This section quantifies the resulting backlog and proposes a mitigation strategy."
)

add_heading_styled("Calculation — Queue Backlog Accumulation", level=2)

add_table(
    ["Parameter", "Symbol", "Value"],
    [
        ["Incoming event rate", "λ (lambda)", "500 events/sec"],
        ["Processing rate", "μ (mu)", "80 events/sec"],
        ["Queue overflow threshold", "Q_max", "10,000 events"],
    ],
)

add_heading_styled("Step 1 — Excess Rate", level=3)
add_body("Excess rate = λ − μ = 500 − 80 = 420 events/sec")
add_body("Result: 420 events per second accumulate in the queue unprocessed.")

add_heading_styled("Step 2 — Time to Reach 10,000 Unprocessed Events", level=3)
add_body("Formula: t = Q_max ÷ Excess rate")
add_body("t = 10,000 ÷ 420 = 23.809... ≈ 23.81 seconds")
add_body("Result: The queue reaches 10,000 unprocessed events in approximately 23.81 seconds.")

add_heading_styled("Step 3 — Verification", level=3)
add_body("Total arrivals at t = 23.81s: λ × t = 500 × 23.81 = 11,905 events")
add_body("Total processed at t = 23.81s: μ × t = 80 × 23.81 = 1,905 events")
add_body("Queue depth = 11,905 − 1,905 = 10,000 events ✓")

add_heading_styled("Bounded Queue Tactic", level=2)
add_body(
    "A bounded queue is a FIFO data structure with a fixed maximum capacity. Once full, any new "
    "enqueue triggers a predefined overflow action. Implementation involves assigning each subscriber "
    "a per-subscriber bounded queue (e.g., Python collections.deque with maxlen). The publish() "
    "method appends to queues rather than calling on_event() directly, and a dedicated worker thread "
    "drains each queue at the subscriber's natural processing rate."
)

add_heading_styled("Eviction Policy — Priority-Based Eviction", level=2)
add_body(
    "This analysis recommends priority-based eviction (drop the least important event) over FIFO "
    "eviction (drop the oldest event). The justification is domain-driven: not all events are equal. "
    "A CongestionAlertEvent with severity CRITICAL represents a public safety hazard, while "
    "VehicleDetectedEvent is routine high-volume telemetry."
)

add_table(
    ["Priority", "Event Type", "Justification"],
    [
        ["1 (highest)", "CongestionAlertEvent", "Public safety; operator action required"],
        ["2", "SpeedViolationEvent", "Legal enforcement; time-sensitive"],
        ["3", "TrafficClearedEvent", "Informational but operationally useful"],
        ["4 (lowest)", "VehicleDetectedEvent", "High-volume routine telemetry"],
    ],
)

add_body(
    "If 90% of the 500 events/sec are vehicle detections, priority eviction reduces the effective "
    "arrival rate for important events to approximately 50 events/sec — well within the "
    "DashboardService's 80 events/sec capacity."
)

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════════
# CLO 4 — SCENARIO 3: DUAL WRITE PROBLEM
# ══════════════════════════════════════════════════════════════════════

add_heading_styled("CLO 4 — Scenario 3: Data Consistency — The Dual Write Problem", level=1)

add_heading_styled("Context", level=2)
add_body(
    "A SpeedViolationEvent is published. AlertService receives it and sends a penalty notice to the "
    "vehicle owner. LoggingService then crashes before persisting the audit log. A penalty now exists "
    "with no record of the violation that triggered it — this is legally inadmissible."
)

add_heading_styled("The Problem — Dual Write Problem", level=2)
add_body(
    "The Dual Write Problem arises when a system must perform two or more writes to independent data "
    "stores as part of a single logical operation, but lacks a shared transactional boundary. In this "
    "system, AlertService (Write 1 — issue penalty) succeeds while LoggingService (Write 2 — persist "
    "audit log) fails. The EventBus delivers envelopes sequentially without a distributed transaction, "
    "so each subscriber operates in its own failure domain. The result is partial completion: an action "
    "was taken without the corresponding record that justifies it."
)

add_table(
    ["Write", "Target", "Outcome"],
    [
        ["Write 1", "AlertService — issue penalty notice", "✓ Succeeded"],
        ["Write 2", "LoggingService — persist audit log", "✗ Failed (crash)"],
    ],
)

add_heading_styled("Solution — The Outbox Pattern", level=2)
add_body(
    "Step 1 — Atomic local write: When a camera detects a speed violation, it performs a single "
    "database transaction that writes both the violation record and a copy of the event payload to "
    "a dedicated outbox table in the same database. Because both writes target the same database, "
    "they are covered by a single ACID transaction."
)
add_body(
    "Step 2 — Background relay: A separate worker polls the outbox table, reads unpublished rows, "
    "constructs EventEnvelopes, and publishes them to the EventBus. It marks rows as published only "
    "after successful delivery."
)
add_body(
    "Step 3 — Retry on failure: If the EventBus is unavailable or a subscriber crashes, the outbox "
    "row remains unpublished. The relay retries on the next polling cycle. The event is never lost."
)

add_table(
    ["Property", "Without Outbox", "With Outbox"],
    [
        ["Event persistence", "In-memory only; lost on crash", "Durable in database"],
        ["Atomicity with source data", "None", "Same-database transaction"],
        ["Retry capability", "None (fire-and-forget)", "Automatic via relay polling"],
        ["At-least-once delivery", "Not guaranteed", "Guaranteed"],
    ],
)

add_heading_styled("Costs of the Outbox Pattern", level=2)
add_body(
    "Added complexity: A new outbox database table, a background relay service, serialization logic, "
    "and operational monitoring. Introduced latency: The relay polls on an interval (100–500ms), adding "
    "a small delay between event recording and subscriber delivery — negligible for traffic enforcement."
)

add_heading_styled("Argument — The Outbox Pattern Is Non-Negotiable", level=2)
add_body(
    "In a traffic enforcement system where penalties must be legally admissible in court, the Outbox "
    "Pattern is not merely beneficial — it is a fundamental requirement. Legal admissibility demands "
    "a complete evidentiary chain: every penalty must be supported by a contemporaneous, tamper-evident "
    "record. A motorist who challenges a fine can present a simple argument: \"There is no system "
    "record showing what violation I committed.\" Without the audit log, the penalty is struck down."
)
add_body(
    "Reliability under failure is a regulatory requirement, not an optimization. Traffic enforcement "
    "systems operate under procurement standards mandating auditability and fault tolerance. The "
    "alternative — hoping LoggingService does not crash — violates every principle of fault-tolerant "
    "design. The Outbox Pattern transforms inevitable component failures from data-loss events into "
    "recoverable delays. The event sits safely in the outbox table, the relay retries, and the audit "
    "trail is complete — perhaps seconds late, but never absent."
)
add_body(
    "The engineering cost of one outbox table and one relay service is bounded and quantifiable. "
    "The cost of unrecorded penalties — mass dismissals, reputational damage, legal liability, and "
    "erosion of public trust — is unbounded. The Outbox Pattern is the architecturally correct, "
    "legally necessary, and ethically responsible choice."
)

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════════
# CONCLUSION
# ══════════════════════════════════════════════════════════════════════

add_heading_styled("Conclusion", level=1)

add_body(
    "This report has presented the complete design, implementation, and critical analysis of an "
    "Event-Driven Traffic Alert System built in Python 3. The system demonstrates three foundational "
    "design patterns — the Observer Pattern for decoupled event delivery via the EventBus, the Event "
    "Envelope Pattern for metadata-rich event wrapping enabling traceability and versioning, and the "
    "Idempotent Receiver Pattern for safe duplicate rejection in AlertService and LoggingService. "
    "The CLO 4 analysis examined three real-world architectural challenges: schema evolution was "
    "addressed through backward-compatible optional fields guided by a formal Architecture Decision "
    "Record; event flooding during peak traffic was mitigated through bounded queues with priority-based "
    "eviction, supported by mathematical verification showing a 23.81-second overflow window; and the "
    "Dual Write Problem was resolved through the Outbox Pattern, ensuring that legally admissible "
    "audit trails are never lost to subscriber failures. Together, these patterns and analyses "
    "demonstrate a system architecture that is extensible, fault-tolerant, and suitable for "
    "safety-critical traffic enforcement in a metropolitan environment."
)

# ── Save ──────────────────────────────────────────────────────────────
output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "report", "CEP_Report.docx")
doc.save(output_path)
print(f"Report generated successfully: {output_path}")
